#!/usr/bin/env python3
"""
Fase 11.0: Preparación y Respaldo Criptográfico DDD
Especificación: .agents/DDD.md
"""

import os
import shutil
import hashlib
import json
import csv
import datetime
import openpyxl

WORKSPACE = "/Users/programacion/Documents/sheets"
TARGET_FILE = os.path.join(WORKSPACE, "Estilo Neutral.xlsx")
BACKUP_DIR = os.path.join(WORKSPACE, "Backups", "2026")
CSV_BACKUP_DIR = os.path.join(BACKUP_DIR, "csv_fase11")

TIMESTAMP_NOW = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=-4)))
TIMESTAMP_STR = TIMESTAMP_NOW.strftime("%Y%m%d_%H%M%S")
TIMESTAMP_ISO = TIMESTAMP_NOW.isoformat()
USER_AUDIT = "Arquitecto DDD Senior (Antigravity Agent)"

def compute_sha256(filepath):
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()

print("════════════════════════════════════════════════════════════")
print("FASE 11.0 — PREPARACIÓN Y BACKUP DDD")
print("════════════════════════════════════════════════════════════")

sha256_inicial = compute_sha256(TARGET_FILE)
file_size = os.path.getsize(TARGET_FILE)
print(f"[11.0.2] SHA-256 Inicial: {sha256_inicial}")
print(f"[11.0.2] Tamaño Inicial:  {file_size} bytes")

os.makedirs(CSV_BACKUP_DIR, exist_ok=True)

# 11.0.1 Backup en 3 formatos
backup_xlsx_name = f"Estilo Neutral_BACKUP_DDD_{TIMESTAMP_STR}.xlsx"
backup_xlsx_path = os.path.join(BACKUP_DIR, backup_xlsx_name)
shutil.copy2(TARGET_FILE, backup_xlsx_path)
print(f"[11.0.1] Respaldo .xlsx creado: {backup_xlsx_path}")

wb = openpyxl.load_workbook(TARGET_FILE, data_only=False)
backup_json = {
    "metadata": {
        "fase": "11.0_ddd_prep",
        "sha256": sha256_inicial,
        "tamano_bytes": file_size,
        "timestamp_iso": TIMESTAMP_ISO,
        "hojas": wb.sheetnames
    },
    "hojas_datos": {}
}

for sname in wb.sheetnames:
    ws = wb[sname]
    safe_name = sname.replace(" ", "_").replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u")
    csv_path = os.path.join(CSV_BACKUP_DIR, f"{safe_name}.csv")
    sdata = []
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as cf:
        writer = csv.writer(cf, quoting=csv.QUOTE_MINIMAL)
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            writer.writerow(["" if v is None else str(v) for v in row_vals])
            sdata.append(row_vals)
    backup_json["hojas_datos"][sname] = sdata

backup_json_path = os.path.join(BACKUP_DIR, f"Estilo Neutral_BACKUP_DDD_{TIMESTAMP_STR}.json")
with open(backup_json_path, "w", encoding="utf-8") as jf:
    json.dump(backup_json, jf, indent=2, ensure_ascii=False)
print(f"[11.0.1] Respaldo JSON estructurado creado: {backup_json_path}")
print(f"[11.0.1] Respaldos CSV (UTF-8 BOM) exportados a: {CSV_BACKUP_DIR}")

# 11.0.3 Confirmar columna validacion
ws_ventas = wb["ventas"]
print(f"[11.0.3] ventas!O1 header: '{ws_ventas['O1'].value}', formula ventas!O2: '{ws_ventas['O2'].value}'")
ws_compras = wb["compras_divisas"]
print(f"[11.0.3] compras_divisas!K1 header: '{ws_compras['K1'].value}', formula compras_divisas!K2: '{ws_compras['K2'].value}'")

print("Fase 11.0 verificada exitosamente.")
